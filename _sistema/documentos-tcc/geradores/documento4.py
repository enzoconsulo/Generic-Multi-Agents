# -*- coding: utf-8 -*-
"""Versao 4: 5 paginas, visual, tabelas enxutas, sem a secao de custo medido."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AQUI = os.path.dirname(os.path.abspath(__file__))
FIGS = os.path.join(AQUI, "figs3")
DESTINO = r"C:\Users\enzoc\OneDrive\Documentos\Gerador_de_projetos\fabrica-multi-agente-arquitetura-4.docx"

AZUL = RGBColor(0x18, 0x4F, 0x95)
TINTA = RGBColor(0x0B, 0x0B, 0x0B)
CINZA = RGBColor(0x52, 0x51, 0x4E)
MUDO = RGBColor(0x89, 0x87, 0x81)
ROXO = RGBColor(0x4A, 0x3A, 0xA7)
VERDE = RGBColor(0x14, 0x86, 0x5D)
VERMELHO = RGBColor(0xB5, 0x2C, 0x2C)
FONTE = "Calibri"
MONO = "Consolas"

doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.left_margin = s.right_margin = Cm(1.8)
s.top_margin = Cm(1.4)
s.bottom_margin = Cm(1.3)


def fonte(run, nome=FONTE, tam=None, cor=None, negrito=None, italico=None):
    run.font.name = nome
    run._element.rPr.rFonts.set(qn("w:eastAsia"), nome)
    if tam:
        run.font.size = Pt(tam)
    if cor is not None:
        run.font.color.rgb = cor
    if negrito is not None:
        run.bold = negrito
    if italico is not None:
        run.italic = italico
    return run


normal = doc.styles["Normal"]
normal.font.name = FONTE
normal.font.size = Pt(9.3)
normal.font.color.rgb = TINTA
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
pf = normal.paragraph_format
pf.space_after = Pt(3)
pf.line_spacing = 1.04
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for nome, tam, cor, antes, depois in [("Heading 1", 12, AZUL, 7, 2),
                                      ("Heading 2", 9.8, TINTA, 4, 1)]:
    st = doc.styles[nome]
    st.font.name = FONTE
    st.font.size = Pt(tam)
    st.font.color.rgb = cor
    st.font.bold = True
    st.font.italic = False
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
    st.paragraph_format.space_before = Pt(antes)
    st.paragraph_format.space_after = Pt(depois)
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


def p(texto="", tam=9.3, cor=TINTA, negrito=False, italico=False,
      alinhamento=None, antes=0, depois=3, espaco=1.04):
    par = doc.add_paragraph()
    if texto:
        fonte(par.add_run(texto), tam=tam, cor=cor, negrito=negrito, italico=italico)
    par.paragraph_format.space_before = Pt(antes)
    par.paragraph_format.space_after = Pt(depois)
    par.paragraph_format.line_spacing = espaco
    if alinhamento is not None:
        par.paragraph_format.alignment = alinhamento
    return par


def rico(pedacos, tam=9.3, depois=3, antes=0):
    par = doc.add_paragraph()
    for pedaco in pedacos:
        t, n = pedaco[0], pedaco[1]
        c = pedaco[2] if len(pedaco) > 2 else TINTA
        fonte(par.add_run(t), tam=tam, cor=c, negrito=n)
    par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    par.paragraph_format.space_after = Pt(depois)
    par.paragraph_format.space_before = Pt(antes)
    par.paragraph_format.line_spacing = 1.04
    return par


def marcador(negrito, resto, tam=9.1):
    par = doc.add_paragraph(style="List Bullet")
    fonte(par.add_run(negrito), tam=tam, negrito=True, cor=TINTA)
    fonte(par.add_run(resto), tam=tam, cor=TINTA)
    par.paragraph_format.space_after = Pt(1)
    par.paragraph_format.line_spacing = 1.02
    par.paragraph_format.left_indent = Cm(0.5)
    par.paragraph_format.first_line_indent = Cm(-0.28)
    par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return par


def figura(arquivo, legenda, largura=13.8):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(2)
    par.paragraph_format.space_after = Pt(1)
    par.add_run().add_picture(os.path.join(FIGS, arquivo), width=Cm(largura))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(4)
    fonte(cap.add_run(legenda), tam=7.4, cor=MUDO, italico=True)


def sombra(celula, cor_hex):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), cor_hex)
    celula._tc.get_or_add_tcPr().append(shd)


def tabela(cabecalho, linhas, larguras, tam=8.2):
    t = doc.add_table(rows=1, cols=len(cabecalho))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D9D9D6")
        borders.append(el)
    t._tbl.tblPr.append(borders)
    for i, texto in enumerate(cabecalho):
        cel = t.rows[0].cells[i]
        cel.text = ""
        par = cel.paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.space_before = Pt(0)
        par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fonte(par.add_run(texto), tam=tam, negrito=True, cor=RGBColor(0xFF, 0xFF, 0xFF))
        sombra(cel, "2A78D6")
    for j, linha in enumerate(linhas):
        cells = t.add_row().cells
        for i, texto in enumerate(linha):
            cells[i].text = ""
            par = cells[i].paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.space_before = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            negrito = texto.startswith("*")
            mono = texto.startswith("`")
            limpo = texto.lstrip("*`")
            fonte(par.add_run(limpo), nome=MONO if mono else FONTE,
                  tam=tam - 0.5 if mono else tam,
                  cor=TINTA if (negrito or mono) else CINZA, negrito=negrito)
            if j % 2 == 1:
                sombra(cells[i], "F5F5F2")
    for linha in t.rows:
        for i, cel in enumerate(linha.cells):
            cel.width = Cm(larguras[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def caixa(titulo, texto, cor="2A78D6", fundo="EAF2FD", cor_titulo=AZUL, tam=8.6):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cel = t.rows[0].cells[0]
    cel.width = Cm(17.4)
    sombra(cel, fundo)
    borders = OxmlElement("w:tblBorders")
    for lado in ("top", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    el = OxmlElement("w:left")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), "18")
    el.set(qn("w:color"), cor)
    borders.append(el)
    t._tbl.tblPr.append(borders)
    par = cel.paragraphs[0]
    par.paragraph_format.space_before = Pt(2)
    par.paragraph_format.space_after = Pt(2)
    par.paragraph_format.line_spacing = 1.02
    par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if titulo:
        fonte(par.add_run(titulo + "  "), tam=tam, negrito=True, cor=cor_titulo)
    fonte(par.add_run(texto), tam=tam, cor=TINTA)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def passos(itens):
    t = doc.add_table(rows=1, cols=len(itens))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    t._tbl.tblPr.append(borders)
    for i, (num, tit, txt) in enumerate(itens):
        cel = t.rows[0].cells[i]
        cel.width = Cm(17.4 / len(itens))
        cel.text = ""
        par = cel.paragraphs[0]
        par.paragraph_format.space_before = Pt(1)
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.line_spacing = 1.02
        fonte(par.add_run(num + "  "), tam=10.5, negrito=True, cor=AZUL)
        fonte(par.add_run(tit), tam=8.2, negrito=True, cor=TINTA)
        par2 = cel.add_paragraph()
        par2.paragraph_format.space_before = Pt(0)
        par2.paragraph_format.space_after = Pt(2)
        par2.paragraph_format.line_spacing = 1.02
        fonte(par2.add_run(txt), tam=7.8, cor=CINZA)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def codigo(linhas, legenda=None):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cel = t.rows[0].cells[0]
    cel.width = Cm(17.4)
    sombra(cel, "F7F7F4")
    borders = OxmlElement("w:tblBorders")
    for lado in ("top", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    el = OxmlElement("w:left")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), "18")
    el.set(qn("w:color"), "C3C2B7")
    borders.append(el)
    t._tbl.tblPr.append(borders)
    primeiro = True
    for l in linhas:
        par = cel.paragraphs[0] if primeiro else cel.add_paragraph()
        par.paragraph_format.space_before = Pt(2 if primeiro else 0)
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.line_spacing = 1.0
        par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        primeiro = False
        comentario = l.strip().startswith("#")
        fonte(par.add_run(l if l else " "), nome=MONO, tam=7.3,
              cor=RGBColor(0x1B, 0xAF, 0x7A) if comentario else TINTA)
    if legenda:
        cap = doc.add_paragraph()
        cap.paragraph_format.space_before = Pt(1)
        cap.paragraph_format.space_after = Pt(4)
        fonte(cap.add_run(legenda), tam=7.4, cor=MUDO, italico=True)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(1)


# ================================================================ CAPA
p("Fábrica de Software Multi-Agente", tam=20, cor=AZUL, negrito=True,
  alinhamento=WD_ALIGN_PARAGRAPH.LEFT, depois=1, espaco=1.0)
p("Como o sistema funciona, e como ele será reconstruído em Elixir/OTP",
  tam=11, cor=CINZA, alinhamento=WD_ALIGN_PARAGRAPH.LEFT, depois=4, espaco=1.0)

reg = doc.add_paragraph()
reg.paragraph_format.space_after = Pt(6)
borda = OxmlElement("w:pBdr")
bot = OxmlElement("w:bottom")
bot.set(qn("w:val"), "single")
bot.set(qn("w:sz"), "12")
bot.set(qn("w:color"), "2A78D6")
borda.append(bot)
reg._p.get_or_add_pPr().append(borda)

p("Trabalho de Conclusão de Curso   ·   Enzo Consulo   ·   agosto de 2026", tam=8.8,
  cor=CINZA, alinhamento=WD_ALIGN_PARAGRAPH.LEFT, depois=6)

caixa("O QUE É",
      "Um sistema que recebe a descrição de um projeto em linguagem natural e o constrói de ponta a "
      "ponta, coordenando agentes especializados que planejam, implementam, verificam, revisam e "
      "documentam — com intervenção humana apenas no pedido inicial. A primeira versão existe e "
      "rodou projetos reais; este documento apresenta o funcionamento do sistema e o projeto da "
      "segunda versão, reconstruída sobre Elixir/OTP. O planejamento detalhado está em documento "
      "à parte.")

# ================================================================ 1
doc.add_heading("1.  O sistema em uma página", level=1)

figura("fig1-fluxo.png", "Figura 1 — Do pedido à entrega. A única intervenção humana obrigatória é "
       "a primeira caixa.")

passos([
    ("1", "Pedido", "O usuário descreve o que quer, uma vez, em linguagem natural."),
    ("2", "Planejamento", "Um agente gera especificação, plano em fases e de 8 a 20 tarefas com "
                          "dependências."),
    ("3", "Construção", "Cada tarefa é implementada por um agente especialista naquela área."),
    ("4", "Dois portões", "Um verifica se funciona; outro revisa se é o que foi pedido. Nenhum "
                          "pode corrigir."),
    ("5", "Entrega", "Tarefa concluída vira commit próprio, e o sistema segue sozinho."),
])

caixa("A IDEIA CENTRAL",
      "Um modelo de linguagem escreve bem em trechos curtos; o que ele não faz sozinho é sustentar "
      "um trabalho longo. As falhas que aparecem aí não são de escrita — são de processo: perda de "
      "contexto, decisão sem rastro, autoaprovação e tentativa sem limite. A resposta é tratar a "
      "construção como linha de produção: tarefas pequenas, estados explícitos, portões operados "
      "por quem não construiu e registro obrigatório a cada transição.")

marcador("Quem implementa nunca é quem aprova. ", "Verificador e revisor não podem corrigir: "
         "reprovam e devolvem.")
marcador("Estado fora da conversa. ", "Se tudo cair agora, a próxima execução reconstrói o mundo "
         "lendo o banco e o git.")
marcador("Falha limitada. ", "Três ciclos por tarefa; no quarto ela é bloqueada e reportada, e a "
         "fila segue.")

# ================================================================ 2
doc.add_heading("2.  O que muda na versão 2", level=1)

p("A primeira versão provou que o desenho funciona. O que ela também mostrou é que quase toda regra "
  "importante do sistema existia como uma frase em português dentro de um prompt — e frase em "
  "prompt é regra opcional: o modelo pode ignorá-la, e às vezes ignora. A versão 2 troca cada uma "
  "dessas frases por um mecanismo que não depende de boa vontade.", depois=3)

tabela(["Na v1 era um pedido no prompt", "Na v2 é uma propriedade da máquina"],
       [["“não encerre com um agente em voo”", "supervisor que percebe a morte do processo"],
        ["“incremente o contador de tentativas”", "coluna do banco, escrita pelo sistema"],
        ["“não escreva fora das suas áreas”", "ferramenta que o processo simplesmente não tem"],
        ["“pare ao estourar o teto de custo”", "processo encerrado por quem o supervisiona"],
        ["“registre o que fez antes de sair”", "transação: grava tudo, ou não grava nada"]],
       [8.7, 8.7])

caixa("A TESE",
      "A segunda versão não precisa ser mais esperta que a primeira — precisa ser mais difícil de "
      "operar errado. O que puder ser um teste, será um teste; o que puder ser uma transação de "
      "banco, será uma transação; o que puder ser um processo supervisionado, não será uma promessa.",
      cor="4A3AA7", fundo="EDEDF8", cor_titulo=ROXO)

# ================================================================ 3
doc.add_heading("3.  A arquitetura", level=1)

p("Cinco camadas, e cada uma só conhece a imediatamente abaixo: quem dispara não sabe qual agente "
  "vai rodar; quem decide a ordem não escreve código; quem escreve código não decide o que vem "
  "depois dele.")

figura("fig3-camadas.png", "Figura 2 — As cinco camadas e a tecnologia de cada uma. A camada de "
       "estado, no rodapé, é a única fonte de verdade.", largura=14.0)

tabela(["Camada", "Escolha", "O que ela compra"],
       [["*Linguagem", "Elixir sobre a BEAM", "processos isolados e supervisão"],
        ["*Painel", "Phoenix + LiveView", "tela ao vivo sem frontend à parte"],
        ["*Banco", "PostgreSQL + Ecto", "transação nas mudanças de estado"],
        ["*Fila", "Oban", "trabalho interrompido volta a ser despachável"],
        ["*Chamada ao modelo", "Req (HTTP direto)", "controle do cache de prompt (seção 5)"],
        ["*Embeddings", "Bumblebee + Nx", "rodam local, sem custo por chamada"],
        ["*Índice semântico", "pgvector", "vetores no mesmo banco das tarefas"],
        ["*Testes", "ExUnit + cliente falso", "suíte sem rede e sem gasto"]],
       [3.4, 4.6, 9.4])

# ================================================================ 4
doc.add_heading("4.  Cada agente é um processo supervisionado", level=1)

p("Esta é a mudança estrutural. Na v1, “agente” era uma sessão passageira que ninguém vigiava, e o "
  "sistema tentava governá-la pedindo bom comportamento no prompt. Em Elixir, agente é um processo: "
  "tem identificador, tem dono, tem quem o encerre e — o que mais importa — tem quem perceba que "
  "ele morreu.")

figura("fig2-supervisao.png", "Figura 3 — A árvore de supervisão. Cada caixa da faixa inferior é um "
       "agente em voo; cada linha da faixa verde é uma regra que hoje é só uma frase escrita.",
       largura=14.0)

marcador("Não existe mais trabalho órfão. ", "Se um agente é cortado no meio, o supervisor percebe "
         "e a tarefa volta para a fila.")
marcador("Paralelismo sem contaminação. ", "Três construtores são três processos isolados; a falha "
         "de um não afeta os outros dois.")
marcador("Teto de custo que corta de verdade. ", "Quem encerra o processo é quem o supervisiona, "
         "não uma instrução dentro do prompt.")

# ================================================================ 5
doc.add_heading("5.  Como um agente trabalha, e como o custo é controlado", level=1)

p("Um agente trabalhando é um laço: monta a requisição, chama a API do modelo e olha o motivo da "
  "parada. Se o modelo pediu ferramentas — ler um arquivo, rodar um teste, escrever código —, o "
  "sistema executa e devolve os resultados; se ele terminou, o laço acaba. Cada volta é uma "
  "requisição paga, e o histórico inteiro é reenviado em todas elas.")

figura("fig4-laco.png", "Figura 4 — O laço de tool use. Os resultados voltam sempre numa única "
       "mensagem, e as ferramentas de uma mesma volta rodam em paralelo.", largura=13.6)

codigo([
    "def handle_info(:trabalhar, estado) do",
    "  case Fabrica.Claude.mensagens(estado.corpo) do",
    "    {:ok, %{stop_reason: \"tool_use\"} = r} ->",
    "      # as ferramentas de uma mesma volta sao independentes: rodam em paralelo",
    "      resultados =",
    "        r.content",
    "        |> Enum.filter(&(&1.type == \"tool_use\"))",
    "        |> Task.async_stream(&Ferramentas.executar(&1, estado.confinamento),",
    "             max_concurrency: 4, timeout: :timer.minutes(2), on_timeout: :kill_task)",
    "        |> Enum.map(&resultado_ou_erro/1)",
    "",
    "      estado",
    "      |> anexar_turno(r.content, resultados)   # TODOS os resultados num turno so",
    "      |> contabilizar(r.usage)                 # custo e tokens gravados por volta",
    "      |> continuar_ou_parar()                  # teto de voltas e de gasto",
    "",
    "    {:ok, %{stop_reason: \"end_turn\"} = r} ->",
    "      {:stop, :normal, finalizar(estado, r)}",
    "  end",
    "end",
], "O núcleo do sistema. O que não aparece aqui — teto de voltas, teto de gasto, prazo, reinício — "
   "não está no prompt: está no supervisor e no estado do processo.")

p("O custo não cresce com o tamanho do texto: cresce com o número de voltas multiplicado pelo "
  "contexto que cada volta carrega. Daí a alavanca principal — a requisição é montada do conteúdo "
  "mais estável para o mais volátil, e marca-se onde ela pode ser reaproveitada.")

figura("fig5-cache.png", "Figura 5 — A ordem dos blocos é a ordem em que o provedor monta a "
       "requisição, e é ela que determina o que pode ser reaproveitado. Ler um trecho já processado "
       "custa um décimo de processá-lo de novo.", largura=14.0)

tabela(["Alavanca de custo", "Efeito"],
       [["Prefixo estável e cacheado", "o trecho repetido cai para um décimo do preço"],
        ["Contexto por papel", "o revisor recebe o diff, não o código inteiro"],
        ["Critérios executáveis rodados antes", "o que falha volta sem gastar um agente"],
        ["Modelo por papel", "verificar é mecânico e usa o modelo barato"],
        ["Trabalho em lote", "metade do preço quando não há pressa"]],
       [6.2, 11.2])

caixa("A ARMADILHA QUE ANULA ISTO, EM SILÊNCIO",
      "O reaproveitamento exige que o trecho estável seja idêntico byte a byte. Uma data, um "
      "contador ou um código de commit dentro dele invalida tudo o que vem depois — sem erro e sem "
      "aviso. Por isso a versão 2 tem um teste que monta o prefixo duas vezes e falha se os bytes "
      "divergirem.",
      cor="E34948", fundo="FDEDED", cor_titulo=VERMELHO)

# ================================================================ 6
doc.add_heading("6.  Memória: o que o sistema sabe, e onde", level=1)

p("Um agente começa sempre sem memória do que houve antes. Por isso a memória do sistema é "
  "deliberada — e, na versão 2, consultável: o agente pede o que precisa em vez de receber arquivos "
  "inteiros.")

figura("fig8-memoria.png", "Figura 6 — As cinco memórias, o que cada uma guarda e quanto custa "
       "lê-la.", largura=14.0)

rico([("O índice denso responde “o que existe e como se chama” — é gerado por script, sem custo de "
       "modelo. O que ele não responde é a pergunta que mais desperdiça ciclo numa fábrica: ", False),
      ("“isto já foi resolvido aqui, e o que foi decidido na época?”", True),
      (". Essa não se responde por nome de função, e sim por significado. O material indexado não é "
       "o código: é a história do projeto — decisões com o motivo registrado, achados de revisão e "
       "tarefas concluídas.", False)])

figura("fig7-rag.png", "Figura 7 — A indexação roda ao commitar, fora do caminho quente; a consulta "
       "roda no despacho, antes de gastar a primeira volta de modelo.", largura=13.8)

caixa("QUANDO NÃO USAR",
      "Busca por significado é aproximada e devolve trechos, não verdades. Ela não substitui o "
      "índice denso, que é exato e grátis, nem responde o que uma consulta comum responde melhor: "
      "“quais tarefas estão bloqueadas” é uma consulta ao banco. Ela entra quando a pergunta é sobre "
      "precedente e o vocabulário exato é desconhecido.",
      cor="1BAF7A", fundo="E7F7F1", cor_titulo=VERDE)

# ================================================================ 7
doc.add_heading("7.  O ciclo de uma tarefa", level=1)

figura("fig6-estados.png", "Figura 8 — Os seis estados. As duas setas vermelhas de volta são os "
       "portões reprovando; a terceira é o limite de três ciclos protegendo a fila.", largura=13.6)

rico([("Quem verifica pergunta “funciona?”", True),
      (" e executa literalmente cada critério de aceite, registrando a evidência. Por isso os "
       "critérios são escritos como comando mais resultado esperado, nunca como frase avaliativa.",
       False)], depois=2)
rico([("Quem revisa pergunta “é o que foi pedido, e está correto?”", True),
      (" — duas checagens distintas. E aqui está o ponto sutil: ", False),
      ("uma entrega pode passar em todos os critérios, não ter defeito nenhum e ainda assim não ser "
       "a tarefa", True),
      (". Critério mal escrito não é licença para entregar outra coisa.", False)])

marcador("Retrabalho sobe de modelo. ", "A partir da segunda tentativa, quando já existe uma "
         "reprovação registrada, o despacho vai para um modelo mais forte.")
marcador("Duas reprovações trocam o especialista. ", "Ele foi escolhido no planejamento, antes de "
         "se saber onde a tarefa iria falhar.")
marcador("Toda transição é uma transação. ", "Novo estado, relatório da etapa e custo entram "
         "juntos, ou não entram.")

p("A mesma fábrica constrói apresentações, documentos e análises. O que separa os dois casos não é "
  "“é código?”, e sim como se prova que ficou pronto: em software a prova vem de graça — o programa "
  "roda e passa ou quebra. Fora dele, a primeira tarefa passa a ser instalar um verificador, e cada "
  "critério é rotulado conforme o grau de prova: executado por comando, inspecionado por script ou "
  "julgado contra itens objetivos.")

# ================================================================ 8
doc.add_heading("8.  Plano de execução e escopo", level=1)

figura("fig12-fases.png", "Figura 9 — As seis fases e os marcos que as encerram. O marco é sempre "
       "um comportamento observável, nunca um percentual de conclusão.", largura=14.0)

p("A ordem não é arbitrária: cada fase entrega algo verificável e destrava a seguinte. A primeira é "
  "a que dá vontade de pular — e é a mais importante, porque sem um cliente falso da API a suíte de "
  "testes custa dinheiro a cada execução, e uma suíte que custa dinheiro deixa de ser executada.",
  depois=3)

tabela(["Dentro do escopo", "Fora do escopo, e por quê"],
       [["Planejar, construir, verificar, revisar e documentar projetos",
         "publicar em produção: custo externo e credencial de terceiros"],
        ["Execução num só computador, com concorrência e supervisão",
         "distribuir em vários: o gargalo é a espera pelo modelo"],
        ["Índice semântico do histórico, com avaliação das buscas",
         "treinar modelos: destruiria a independência de fornecedor"],
        ["Painel local para acompanhar e disparar o trabalho",
         "vários usuários e autenticação: é ferramenta de um operador"]],
       [8.7, 8.7])

doc.add_heading("Limites conhecidos do desenho", level=2)

tabela(["Limite", "Como o desenho responde"],
       [["O ganho de cache pode não se confirmar",
         "teste que monta o prefixo duas vezes e compara byte a byte"],
        ["A busca semântica pode trazer trecho inútil",
         "conjunto de perguntas com resposta conhecida, medido a cada mudança"],
        ["Uma só árvore de arquivos por projeto",
         "as áreas declaradas viram exclusão mútua verificada pelo motor"],
        ["O portão do meio é frágil fora de software",
         "grau de prova obrigatório, com a proporção visível no painel"]],
       [6.2, 11.2])

rico([("O que o trabalho pretende demonstrar: ", True),
      ("que a diferença entre um assistente de programação e uma linha de produção está inteiramente "
       "na camada de governança — papéis com autoridade explícita, estado transacional, portões "
       "operados por quem não construiu e falha limitada por desenho. Nada disso depende de um "
       "modelo específico ou de um fornecedor.", False)])

doc.save(DESTINO)
print("gerado:", DESTINO)
