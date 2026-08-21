# -*- coding: utf-8 -*-
"""Base do documento v7 (COMPLETO): estilos e helpers. O conteudo vem dos corpo7*.py."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AQUI = os.path.dirname(os.path.abspath(__file__))
FIGS = os.path.join(AQUI, "figs3")
DESTINO = r"C:\Users\enzoc\OneDrive\Documentos\Gerador_de_projetos\fabrica-multi-agente-arquitetura-7.docx"

AZUL = RGBColor(0x18, 0x4F, 0x95)
AZUL_CLARO = RGBColor(0x2A, 0x78, 0xD6)
TINTA = RGBColor(0x0B, 0x0B, 0x0B)
CINZA = RGBColor(0x52, 0x51, 0x4E)
MUDO = RGBColor(0x89, 0x87, 0x81)
ROXO = RGBColor(0x4A, 0x3A, 0xA7)
VERDE = RGBColor(0x14, 0x86, 0x5D)
LARANJA = RGBColor(0xC2, 0x4E, 0x1E)
VERMELHO = RGBColor(0xB5, 0x2C, 0x2C)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
FONTE = "Calibri"
MONO = "Consolas"

doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.left_margin = s.right_margin = Cm(1.9)
s.top_margin = Cm(1.5)
s.bottom_margin = Cm(1.4)


def fonte(run, nome=FONTE, tam=None, cor=None, negrito=None, italico=None):
    run.font.name = nome
    run._element.rPr.rFonts.set(qn("w:eastAsia"), nome)
    if tam:
        run.font.size = Pt(tam)
    if cor is not None:
        run.font.color.rgb = cor
    if negrito is not None:
        run.bold = negrito
    if italico is not None:
        run.italic = italico
    return run


normal = doc.styles["Normal"]
normal.font.name = FONTE
normal.font.size = Pt(9.5)
normal.font.color.rgb = TINTA
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
pf = normal.paragraph_format
pf.space_after = Pt(3)
pf.line_spacing = 1.05
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for nome, tam, cor, antes, depois in [("Heading 1", 12.5, AZUL, 8, 2),
                                      ("Heading 2", 10, TINTA, 5, 1)]:
    st = doc.styles[nome]
    st.font.name = FONTE
    st.font.size = Pt(tam)
    st.font.color.rgb = cor
    st.font.bold = True
    st.font.italic = False
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
    st.paragraph_format.space_before = Pt(antes)
    st.paragraph_format.space_after = Pt(depois)
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


def p(texto="", tam=9.5, cor=TINTA, negrito=False, italico=False,
      alinhamento=None, antes=0, depois=3, espaco=1.05):
    par = doc.add_paragraph()
    if texto:
        fonte(par.add_run(texto), tam=tam, cor=cor, negrito=negrito, italico=italico)
    par.paragraph_format.space_before = Pt(antes)
    par.paragraph_format.space_after = Pt(depois)
    par.paragraph_format.line_spacing = espaco
    if alinhamento is not None:
        par.paragraph_format.alignment = alinhamento
    return par


def rico(pedacos, tam=9.5, depois=3, antes=0):
    par = doc.add_paragraph()
    for pedaco in pedacos:
        t, n = pedaco[0], pedaco[1]
        c = pedaco[2] if len(pedaco) > 2 else TINTA
        fonte(par.add_run(t), tam=tam, cor=c, negrito=n)
    par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    par.paragraph_format.space_after = Pt(depois)
    par.paragraph_format.space_before = Pt(antes)
    par.paragraph_format.line_spacing = 1.05
    return par


def marcador(negrito, resto, tam=9.3):
    par = doc.add_paragraph(style="List Bullet")
    fonte(par.add_run(negrito), tam=tam, negrito=True, cor=TINTA)
    fonte(par.add_run(resto), tam=tam, cor=TINTA)
    par.paragraph_format.space_after = Pt(1)
    par.paragraph_format.line_spacing = 1.03
    par.paragraph_format.left_indent = Cm(0.5)
    par.paragraph_format.first_line_indent = Cm(-0.28)
    par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return par


def figura(arquivo, legenda, largura=14.0):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(3)
    par.paragraph_format.space_after = Pt(1)
    par.add_run().add_picture(os.path.join(FIGS, arquivo), width=Cm(largura))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(5)
    fonte(cap.add_run(legenda), tam=7.5, cor=MUDO, italico=True)


def sombra(celula, cor_hex):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), cor_hex)
    celula._tc.get_or_add_tcPr().append(shd)


def _sem_bordas(t, lados=("top", "left", "bottom", "right", "insideH", "insideV")):
    borders = OxmlElement("w:tblBorders")
    for lado in lados:
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    t._tbl.tblPr.append(borders)


def tabela(cabecalho, linhas, larguras, tam=8.2):
    t = doc.add_table(rows=1, cols=len(cabecalho))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D9D9D6")
        borders.append(el)
    t._tbl.tblPr.append(borders)
    for i, texto in enumerate(cabecalho):
        cel = t.rows[0].cells[i]
        cel.text = ""
        par = cel.paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.space_before = Pt(0)
        par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fonte(par.add_run(texto), tam=tam, negrito=True, cor=BRANCO)
        sombra(cel, "2A78D6")
    for j, linha in enumerate(linhas):
        cells = t.add_row().cells
        for i, texto in enumerate(linha):
            cells[i].text = ""
            par = cells[i].paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.space_before = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            negrito = texto.startswith("*")
            mono = texto.startswith("`")
            limpo = texto.lstrip("*`")
            fonte(par.add_run(limpo), nome=MONO if mono else FONTE,
                  tam=tam - 0.5 if mono else tam,
                  cor=TINTA if (negrito or mono) else CINZA, negrito=negrito)
            if j % 2 == 1:
                sombra(cells[i], "F5F5F2")
    for linha in t.rows:
        for i, cel in enumerate(linha.cells):
            cel.width = Cm(larguras[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def caixa(titulo, texto, cor="2A78D6", fundo="EAF2FD", cor_titulo=AZUL, tam=8.7):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cel = t.rows[0].cells[0]
    cel.width = Cm(17.4)
    sombra(cel, fundo)
    _sem_bordas(t, ("top", "bottom", "right", "insideH", "insideV"))
    el = OxmlElement("w:left")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), "18")
    el.set(qn("w:color"), cor)
    t._tbl.tblPr[-1].append(el)
    par = cel.paragraphs[0]
    par.paragraph_format.space_before = Pt(2)
    par.paragraph_format.space_after = Pt(2)
    par.paragraph_format.line_spacing = 1.03
    par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if titulo:
        fonte(par.add_run(titulo + "  "), tam=tam, negrito=True, cor=cor_titulo)
    fonte(par.add_run(texto), tam=tam, cor=TINTA)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def parte(numero, titulo_txt, resumo, cor="2A78D6"):
    """Divisor de parte: faixa colorida com numero, titulo e uma linha de resumo."""
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cel = t.rows[0].cells[0]
    cel.width = Cm(17.4)
    sombra(cel, cor)
    _sem_bordas(t)
    par = cel.paragraphs[0]
    par.paragraph_format.space_before = Pt(3)
    par.paragraph_format.space_after = Pt(1)
    par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fonte(par.add_run("PARTE " + numero + "   "), tam=8.5, negrito=True, cor=BRANCO)
    fonte(par.add_run(titulo_txt), tam=12.5, negrito=True, cor=BRANCO)
    par2 = cel.add_paragraph()
    par2.paragraph_format.space_before = Pt(0)
    par2.paragraph_format.space_after = Pt(3)
    par2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fonte(par2.add_run(resumo), tam=8.6, cor=BRANCO, italico=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def passos(itens):
    t = doc.add_table(rows=1, cols=len(itens))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _sem_bordas(t)
    for i, (num, tit, txt) in enumerate(itens):
        cel = t.rows[0].cells[i]
        cel.width = Cm(17.4 / len(itens))
        cel.text = ""
        par = cel.paragraphs[0]
        par.paragraph_format.space_before = Pt(1)
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.line_spacing = 1.02
        fonte(par.add_run(num + "  "), tam=10.5, negrito=True, cor=AZUL)
        fonte(par.add_run(tit), tam=8.2, negrito=True, cor=TINTA)
        par2 = cel.add_paragraph()
        par2.paragraph_format.space_before = Pt(0)
        par2.paragraph_format.space_after = Pt(2)
        par2.paragraph_format.line_spacing = 1.02
        fonte(par2.add_run(txt), tam=7.8, cor=CINZA)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def codigo(linhas, legenda=None):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cel = t.rows[0].cells[0]
    cel.width = Cm(17.4)
    sombra(cel, "F7F7F4")
    _sem_bordas(t, ("top", "bottom", "right", "insideH", "insideV"))
    el = OxmlElement("w:left")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), "18")
    el.set(qn("w:color"), "C3C2B7")
    t._tbl.tblPr[-1].append(el)
    primeiro = True
    for l in linhas:
        par = cel.paragraphs[0] if primeiro else cel.add_paragraph()
        par.paragraph_format.space_before = Pt(2 if primeiro else 0)
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.line_spacing = 1.0
        par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        primeiro = False
        comentario = l.strip().startswith("#") or l.strip().startswith("--")
        fonte(par.add_run(l if l else " "), nome=MONO, tam=7.4,
              cor=RGBColor(0x1B, 0xAF, 0x7A) if comentario else TINTA)
    if legenda:
        cap = doc.add_paragraph()
        cap.paragraph_format.space_before = Pt(1)
        cap.paragraph_format.space_after = Pt(5)
        fonte(cap.add_run(legenda), tam=7.5, cor=MUDO, italico=True)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(1)


def quebra():
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(0)
    par.add_run().add_break(WD_BREAK.PAGE)
