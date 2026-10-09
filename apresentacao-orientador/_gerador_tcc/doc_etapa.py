# -*- coding: utf-8 -*-
"""Documento de apresentação de UMA etapa (Word, paisagem, uma ideia por página).

    python doc_etapa.py <n> <saida.docx>

Toda página traz no topo a faixa "Etapa N de 10" com a linha do tempo. Conteúdo: capa (proposta ×
resultado), onde a etapa entra no sistema, como funciona (diagrama), cada tarefa em cartão, a prova
do marco passo a passo, números medidos na branch, requisitos e próxima etapa.
"""
import json
import os
import re
import sys
import tempfile

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.section import WD_ORIENT  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
import catalogo as C  # noqa: E402
import figuras_etapa as F  # noqa: E402
import metricas as M  # noqa: E402
from dados_deck import REQ_STATUS, REQ_TXT  # noqa: E402
from textos_etapas import TEXTOS  # noqa: E402
from tickets_e1 import E1  # noqa: E402

DET = dict(E1)
if os.path.exists(os.path.join(AQUI, "tickets_detalhe.json")):
    DET.update(json.load(open(os.path.join(AQUI, "tickets_detalhe.json"), encoding="utf-8")))

NAVY, OKC, WARNC, BADC, CINZA, AZUL = "1B2A41", "2A9D8F", "E9A23B", "C0392B", "F2F4F7", "4A6FA5"
RGB = lambda h: RGBColor(int(h[:2], 16), int(h[2:4], 16), int(h[4:], 16))  # noqa: E731
STATUS_TXT = {"done": "Concluída", "blocked": "Bloqueada", "in_progress": "Em andamento", "todo": "A fazer"}
STATUS_COR = {"done": OKC, "blocked": BADC, "in_progress": WARNC, "todo": "8C96A3"}
ETAPA_STATUS = {1: "parcial", 2: "ok", 3: "ok", 4: "ok", 5: "ok", 6: "bloqueada", 7: "andamento"}
ETAPA_TXT = {"ok": ("ENTREGA APROVADA", OKC), "parcial": ("ENTREGUE COM PENDÊNCIA", WARNC),
             "bloqueada": ("ENTREGA BLOQUEADA", BADC), "andamento": ("ENTREGA EM ANDAMENTO", WARNC)}
LARG = 26.5  # largura útil em cm (A4 paisagem, margens de 1,6 cm)


# ---------------------------------------------------------------- utilitários de formatação
def sombra(cel, cor):
    tc = cel._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), cor)
    tc.append(sh)


def sem_bordas(t):
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "nil")
        b.append(e)
    t._tbl.tblPr.append(b)


def margem_celula(cel, cm=0.3):
    tcPr = cel._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for lado in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:w"), str(int(cm * 567)))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tcPr.append(mar)


def run(p, texto, tam=12, cor="1F2328", negrito=False, italico=False, fonte=None):
    r = p.add_run(texto)
    r.font.size = Pt(tam)
    r.font.color.rgb = RGB(cor)
    r.bold = negrito
    r.italic = italico
    if fonte:
        r.font.name = fonte
    return r


def run_md(p, texto, tam=12, cor="1F2328", negrito=False):
    """Escreve texto com `trechos de código` em Consolas, sem as crases."""
    for i, pedaco in enumerate(texto.split("`")):
        if pedaco:
            run(p, pedaco, tam - 0.5 if i % 2 else tam, "2A4A7F" if i % 2 else cor, negrito, fonte="Consolas" if i % 2 else None)


def nao_quebra(tabela_):
    for row in tabela_.rows:
        trPr = row._tr.get_or_add_trPr()
        e = OxmlElement("w:cantSplit")
        trPr.append(e)
        for cel in row.cells:
            for p in cel.paragraphs:
                p.paragraph_format.keep_with_next = True


def par(alvo, texto="", tam=12, cor="1F2328", negrito=False, antes=0, depois=4, alinhar=None, fonte=None):
    p = alvo.add_paragraph()
    p.paragraph_format.space_before = Pt(antes)
    p.paragraph_format.space_after = Pt(depois)
    if alinhar is not None:
        p.alignment = alinhar
    if texto:
        run(p, texto, tam, cor, negrito, fonte=fonte)
    return p


def tabela(doc, linhas, colunas, larguras):
    t = doc.add_table(rows=linhas, cols=colunas)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    sem_bordas(t)
    for i, w in enumerate(larguras):
        t.columns[i].width = Cm(w)
        for r in t.rows:
            r.cells[i].width = Cm(w)
    return t


def limpa(cel):
    p = cel.paragraphs[0]._p
    p.getparent().remove(p)


def pronta(cel, fundo, margem=0.35):
    sombra(cel, fundo)
    margem_celula(cel, margem)
    limpa(cel)
    return cel


def titulo_pagina(doc, texto, sub=None):
    p = par(doc, texto, 24, NAVY, True, depois=2, fonte="Cambria")
    p.paragraph_format.page_break_before = True
    if sub:
        par(doc, sub, 12.5, "6B7480", depois=10)


def caixa(doc, rotulo, texto, fundo, tam=15):
    c = pronta(tabela(doc, 1, 1, [LARG]).rows[0].cells[0], fundo, 0.45)
    if rotulo:
        par(c, rotulo, 10.5, "FFFFFF", True, depois=4)
    par(c, texto, tam, "FFFFFF", True, depois=0, fonte="Cambria")


# ---------------------------------------------------------------- páginas
def capa(doc, n, e, tx, tarefas, cod):
    rot, cor = ETAPA_TXT[ETAPA_STATUS[n]]
    par(doc, f"ETAPA {n} DE 10", 13, "6B7480", True, antes=4, depois=0)
    par(doc, e["pt"], 34, NAVY, True, depois=2, fonte="Cambria")
    p = par(doc, depois=14)
    run(p, f"Período planejado: {e['datas']}     ·     branch ", 12, "6B7480")
    run(p, f"etapa/{e['slug']}", 12, AZUL, True, fonte="Consolas")
    a, b = tabela(doc, 1, 2, [13.0, 13.5]).rows[0].cells
    pronta(a, NAVY, 0.5)
    pronta(b, cor, 0.5)
    par(a, "PROPOSTA DO PLANO", 11, WARNC, True, depois=6)
    par(a, e["entrega_pt"], 19, "FFFFFF", True, depois=0, fonte="Cambria")
    par(b, rot, 11, "FFFFFF", True, depois=6)
    par(b, tx["resumo"], 15, "FFFFFF", True, depois=0, fonte="Cambria")
    par(doc, depois=10)
    concl = sum(1 for x in tarefas if x["status"] == "done")
    if cod:
        nums = [(f"{concl}/{len(tarefas)}", "tarefas concluídas"), (str(len(cod["commits"])), "commits nesta etapa"),
                (f"{cod['testes']:,}".replace(",", "."), "testes automáticos"), (str(cod["modulos"]), "módulos Elixir")]
    else:
        nums = [(f"{concl}/{len(tarefas)}", "documentos"), ("15", "requisitos"), ("8", "critérios de desempenho"), ("5", "projetos de teste")]
    for c, (v, r) in zip(tabela(doc, 1, 4, [6.6] * 4).rows[0].cells, nums):
        pronta(c, CINZA)
        par(c, v, 30, NAVY, True, depois=0, fonte="Cambria")
        par(c, r, 11, "6B7480", depois=0)
    par(doc, depois=12)
    par(doc, "NESTA APRESENTAÇÃO", 10.5, AZUL, True, depois=4)
    roteiro = ["Onde entra no sistema", "Como funciona", "Tarefa por tarefa", "Como se prova"] + (["Números"] if cod else []) + ["Próximos passos"]
    for c, (i, nome) in zip(tabela(doc, 1, len(roteiro), [LARG / len(roteiro)] * len(roteiro)).rows[0].cells, enumerate(roteiro, 1)):
        pronta(c, "FFFFFF", 0.15)
        p = par(c, depois=0)
        run(p, f"{i}  ", 14, WARNC, True)
        run(p, nome, 11.5, NAVY, True)


def pagina_mapa(doc, figs, n=2):
    sub = ("A E1 não constrói camada: ela desenha o sistema inteiro. A construção começa pela base, na E2, e sobe uma camada por etapa."
           if n == 1 else "O sistema é construído de baixo para cima; cada etapa acrescenta uma camada. Em destaque, a camada desta etapa.")
    titulo_pagina(doc, "Onde esta etapa entra no sistema", sub)
    doc.add_picture(figs["mapa"], width=Cm(21))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def pagina_como(doc, tx, figs):
    titulo_pagina(doc, "Como funciona", "O desenho mostra as peças desta etapa e como elas conversam.")
    doc.add_picture(figs["diag"], width=Cm(23.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for c in tx["como"]:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        run_md(p, c, 12)


def cartao(doc, t):
    d = DET.get(t["key"], {})
    dp = d.get("pt", {})
    st = t["status"]
    tb = tabela(doc, 2, 2, [20.5, 6.0])
    h1, h2 = tb.rows[0].cells
    pronta(h1, NAVY, 0.28)
    pronta(h2, STATUS_COR[st], 0.28)
    p = par(h1, depois=0)
    run(p, t["key"] + "   ", 12, WARNC, True, fonte="Consolas")
    run(p, t["pt"]["titulo"], 13.5, "FFFFFF", True)
    par(h2, ("✓ " if st == "done" else "■ ") + STATUS_TXT[st], 12.5, "FFFFFF", True, depois=0, alinhar=WD_ALIGN_PARAGRAPH.CENTER)
    corpo = pronta(tb.rows[1].cells[0].merge(tb.rows[1].cells[1]), CINZA)
    p = par(corpo, depois=4)
    run(p, "O que é   ", 11, AZUL, True)
    run_md(p, dp.get("resumo") or t["pt"]["desc"], 11.5)
    if dp.get("porque"):
        p = par(corpo, depois=4)
        run(p, "Por que   ", 11, AZUL, True)
        run_md(p, dp["porque"], 11.5)
    for c in (dp.get("como") or [])[:4]:
        p = corpo.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        run_md(p, c, 10.5)
    ver = [v.strip("`") for v in d.get("verificar", [])]
    if ver:
        p = par(corpo, antes=4, depois=0)
        run(p, "Verificar   ", 11, AZUL, True)
        run(p, "    ".join(ver), 9.5, fonte="Consolas")
    nao_quebra(tb)
    par(doc, depois=6)


def pagina_tarefas(doc, tarefas):
    titulo_pagina(doc, "O que foi feito, tarefa por tarefa",
                  "Cada tarefa é um pacote do plano de desenvolvimento. Detalhes completos em docs/tasks/pt/.")
    for t in tarefas:
        if t["tipo"] != "Milestone":
            cartao(doc, t)


def pagina_prova(doc, n, tarefas, tx, cod):
    marco = next((t for t in tarefas if t["tipo"] == "Milestone"), None)
    titulo_pagina(doc, "Como se prova a entrega", "A entrega combinada no plano, demonstrada por um teste ou execução real.")
    caixa(doc, "ENTREGA COMBINADA", C.ETAPAS[n]["entrega_pt"], NAVY, tam=16)
    par(doc, depois=8)
    if marco:
        d = DET.get(marco["key"], {})
        passos = d.get("pt", {}).get("como") or []
        if marco["status"] == "done" and passos:
            par(doc, "O que o teste faz, passo a passo", 14, NAVY, True, depois=4)
            for i, s in enumerate(passos, 1):
                p = par(doc, depois=3)
                run(p, f"{i}.  ", 13, WARNC, True)
                run_md(p, re.sub(r"^\s*\d+[.)]\s*", "", s), 12)
        elif marco["status"] != "done":
            caixa(doc, "SITUAÇÃO DO MARCO", d.get("pt", {}).get("resumo", ""), BADC if marco["status"] == "blocked" else WARNC, tam=13)
    par(doc, depois=6)
    c = pronta(tabela(doc, 1, 1, [LARG]).rows[0].cells[0], "EEF2F7")
    par(c, "COMO DEMONSTRAR", 10.5, AZUL, True, depois=3)
    par(c, tx["prova"], 12, depois=0, fonte="Consolas")
    if cod:
        par(doc, f"Na branch desta etapa, a verificação completa (mix factory.ci) passa nos 6 estágios, com {cod['testes']} testes automáticos.",
            11.5, "6B7480", antes=6)


def grafico(n, cods, caminho):
    etapas = [k for k in range(2, n + 1) if cods.get(k)]
    fig, ax = plt.subplots(1, 2, figsize=(11, 3.2), dpi=180)
    for a, vals, tit, fmt in [(ax[0], [cods[k]["testes"] for k in etapas], "Testes automáticos", "{:.0f}"),
                              (ax[1], [cods[k]["linhas_lib"] / 1000 for k in etapas], "Código (mil linhas em lib/)", "{:.1f}")]:
        b = a.bar([f"E{k}" for k in etapas], vals, color=["#2a78d6" if k == n else "#a9c6ee" for k in etapas], width=0.6)
        a.set_title(tit, fontsize=11, loc="left", color="#333")
        a.spines[["top", "right", "left"]].set_visible(False)
        a.tick_params(axis="y", left=False, labelleft=False)
        for r, v in zip(b, vals):
            a.text(r.get_x() + r.get_width() / 2, r.get_height(), fmt.format(v), ha="center", va="bottom", fontsize=9)
    fig.tight_layout()
    fig.savefig(caminho)
    plt.close(fig)


def pagina_numeros(doc, n, tarefas, cods, proc):
    cod, ant = cods.get(n), cods.get(n - 1)
    titulo_pagina(doc, "Números da etapa", "Medidos no código da branch desta etapa; o crescimento compara com a etapa anterior.")

    def delta(k):
        return f"+{cod[k] - ant[k]:,}".replace(",", ".") + " nesta etapa" if ant else "base do sistema"
    reais = [proc[t["key"]] for t in tarefas if t["key"] in proc]
    ex = sum(x["execucoes"] for x in reais) / (sum(x["subtarefas"] for x in reais) or 1)
    blocos = [(f"{cod['testes']:,}".replace(",", "."), "testes automáticos", delta("testes")),
              (f"{cod['linhas_lib']:,}".replace(",", "."), "linhas em lib/", delta("linhas_lib")),
              (str(cod["modulos"]), "módulos Elixir", delta("modulos")),
              (f"{ex:.2f}".replace(".", ","), "execuções por item", "critério: até 2")]
    for c, (v, r, dl) in zip(tabela(doc, 1, 4, [6.6] * 4).rows[0].cells, blocos):
        pronta(c, CINZA)
        par(c, v, 28, NAVY, True, depois=0, fonte="Cambria")
        par(c, r, 11, "6B7480", depois=0)
        par(c, dl, 10.5, OKC, True, depois=0)
    png = os.path.join(tempfile.gettempdir(), f"cresc-e{n}.png")
    grafico(n, cods, png)
    par(doc, depois=6)
    doc.add_picture(png, width=Cm(23))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    par(doc, "Execuções por item: quantas vezes um agente precisou construir cada item até ele ser aprovado pelos dois portões.",
        10.5, "6B7480", antes=2)


def pagina_final(doc, n, tarefas, tx):
    titulo_pagina(doc, "Requisitos, pendências e próxima etapa")
    reqs = sorted({r for t in tarefas for r in t["req"]}, key=lambda x: int(x[1:]))
    if reqs:
        par(doc, "Requisitos atendidos nesta etapa", 14, NAVY, True, depois=4)
        t = tabela(doc, len(reqs), 2, [2.2, 24.3])
        for row, r in zip(t.rows, reqs):
            s, como = REQ_STATUS[r]
            a, b = row.cells
            pronta(a, {"ok": OKC, "parcial": WARNC, "pendente": BADC}[s], 0.2)
            pronta(b, "FFFFFF", 0.2)
            par(a, r, 13, "FFFFFF", True, depois=0, alinhar=WD_ALIGN_PARAGRAPH.CENTER)
            p = par(b, depois=0)
            run(p, REQ_TXT[r] + "   ", 12.5, NAVY, True)
            run(p, {"ok": "atendido", "parcial": "parcial", "pendente": "pendente"}[s] + " — " + como, 11, "6B7480")
        par(doc, depois=10)
    par(doc, "O que ainda falta", 14, NAVY, True, depois=4)
    for x in tx["pendente"] or ["Nada pendente nesta etapa."]:
        p = doc.add_paragraph(style="List Bullet")
        run(p, x, 12)
    par(doc, depois=10)
    caixa(doc, "PRÓXIMA ETAPA", tx["proxima"], OKC if ETAPA_STATUS[n] == "ok" else NAVY, tam=15)


def apendice(doc, cod):
    if not cod or not cod["commits"]:
        return
    titulo_pagina(doc, "Commits desta etapa", "Na ordem em que entraram na branch.")
    for _, s in cod["commits"]:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(0)
        run(p, s, 10, fonte="Consolas")


def gerar(n, saida):
    e = C.ETAPAS[n]
    tx = TEXTOS[n]
    tarefas = [t for t in C.T if t["etapa"] == n]
    cods = {k: M.codigo(k) for k in range(2, n + 1)}
    cod = cods.get(n)
    proc = M.processo()
    d = tempfile.mkdtemp()
    figs = {k: os.path.join(d, k + ".png") for k in ("faixa", "mapa", "diag")}
    F.faixa(n, ETAPA_STATUS, figs["faixa"])
    F.mapa(n, figs["mapa"])
    F.diagrama(n, figs["diag"])

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.6)
    sec.top_margin = Cm(3.3)
    sec.bottom_margin = Cm(1.4)
    sec.header_distance = Cm(0.6)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(12)
    cab = sec.header.paragraphs[0]
    cab.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cab.add_run().add_picture(figs["faixa"], width=Cm(26.5))
    run(sec.footer.paragraphs[0], f"Fábrica de Software Multi-Agente · TCC · Enzo Consulo   —   Etapa {n}: {e['pt']}", 9, "8C96A3")

    capa(doc, n, e, tx, tarefas, cod)
    pagina_mapa(doc, figs, n)
    pagina_como(doc, tx, figs)
    pagina_tarefas(doc, tarefas)
    pagina_prova(doc, n, tarefas, tx, cod)
    if cod:
        pagina_numeros(doc, n, tarefas, cods, proc)
    pagina_final(doc, n, tarefas, tx)
    apendice(doc, cod)
    os.makedirs(os.path.dirname(os.path.abspath(saida)), exist_ok=True)
    doc.save(saida)
    print("ok", saida)


if __name__ == "__main__":
    gerar(int(sys.argv[1]), sys.argv[2])
